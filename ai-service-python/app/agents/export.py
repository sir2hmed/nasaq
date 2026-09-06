"""Generate genuine Markdown, PDF, and DOCX artifacts on local storage."""

import hashlib
import re
from html import escape
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from app.agents.base import AgentExecution, BaseAgent, validation_error
from app.domain import AgentArtifact, ExecutionContext, ExecutionError

MIME_TYPES = {
    "markdown": "text/markdown; charset=utf-8",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
EXTENSIONS = {"markdown": "md", "pdf": "pdf", "docx": "docx"}


class ExportAgent(BaseAgent):
    agent_type = "export"
    input_schema = "Writer structured text output or Video scene plan"
    output_schema = "real file artifacts with MIME type, size, and checksum"

    def __init__(self, artifact_root: Path) -> None:
        self.artifact_root = artifact_root

    def validate_config(self, config: dict[str, Any], node_id: str) -> list[ExecutionError]:
        formats = config.get("formats")
        if not isinstance(formats, list) or not formats:
            return [
                validation_error(
                    "invalid_config",
                    "Export requires at least one output format.",
                    node_id,
                    self.agent_type,
                    field="formats",
                )
            ]
        unsupported = [value for value in formats if value not in MIME_TYPES]
        if unsupported:
            return [
                validation_error(
                    "invalid_config",
                    "Export contains an unsupported output format.",
                    node_id,
                    self.agent_type,
                    field="formats",
                    unsupported=unsupported,
                    allowed=sorted(MIME_TYPES),
                )
            ]
        if len(set(formats)) != len(formats):
            return [
                validation_error(
                    "invalid_config",
                    "Export formats must not contain duplicates.",
                    node_id,
                    self.agent_type,
                    field="formats",
                )
            ]
        return []

    def validate_input(self, input_data: dict[str, Any], node_id: str) -> list[ExecutionError]:
        if self._written_content(input_data) is None:
            return [
                validation_error(
                    "invalid_input",
                    "Export requires structured Writer output or a Video scene plan.",
                    node_id,
                    self.agent_type,
                    expected="text",
                )
            ]
        return []

    def execute(
        self,
        input_data: dict[str, Any],
        config: dict[str, Any],
        execution_context: ExecutionContext,
        node_id: str,
    ) -> AgentExecution:
        written = self._written_content(input_data)
        assert written is not None
        output_directory = self._output_directory(execution_context.workflow_run_id, node_id)
        output_directory.mkdir(parents=True, exist_ok=True)

        artifacts: list[AgentArtifact] = []
        for output_format in config["formats"]:
            path = output_directory / f"demo-output.{EXTENSIONS[output_format]}"
            if output_format == "markdown":
                self._write_markdown(path, written)
            elif output_format == "pdf":
                self._write_pdf(path, written)
            else:
                self._write_docx(path, written)
            artifacts.append(self._artifact(path, output_format, execution_context, node_id))

        return AgentExecution(
            output_type="files",
            data={
                "files": [artifact.model_dump(mode="json") for artifact in artifacts],
                "formats": config["formats"],
                "demo_notice": (
                    "DEMO MODE — these are real local files generated from demo content."
                ),
            },
            artifacts=artifacts,
            warnings=["DEMO MODE: files contain simulated research and writing output."],
        )

    @staticmethod
    def _written_content(input_data: dict[str, Any]) -> dict[str, Any] | None:
        for value in input_data.values():
            required = {"title", "content", "language", "format"}
            if isinstance(value, dict) and required <= value.keys():
                return value
            video_required = {"title", "scenes", "language", "video"}
            if isinstance(value, dict) and video_required <= value.keys():
                scene_text = "\n\n".join(
                    f"## {scene.get('title', 'Scene')}\n\n{scene.get('text', '')}"
                    for scene in value["scenes"]
                    if isinstance(scene, dict)
                )
                return {
                    "title": f"{value['title']} — Video production notes",
                    "content": scene_text,
                    "language": value["language"],
                    "format": "video_notes",
                }
        return None

    def _output_directory(self, run_id: str, node_id: str) -> Path:
        safe_run = self._safe_component(run_id)
        safe_node = self._safe_component(node_id)
        return self.artifact_root.resolve() / safe_run / safe_node

    @staticmethod
    def _safe_component(value: str) -> str:
        cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", value).strip(".-")
        return cleaned[:100] or "unknown"

    @staticmethod
    def _write_markdown(path: Path, written: dict[str, Any]) -> None:
        document = (
            f"# {written['title']}\n\n"
            "> DEMO MODE — generated from simulated research; not for factual reliance.\n\n"
            f"{written['content']}\n"
        )
        path.write_text(document, encoding="utf-8", newline="\n")

    @classmethod
    def _write_pdf(cls, path: Path, written: dict[str, Any]) -> None:
        font_name = cls._unicode_font()
        styles = getSampleStyleSheet()
        body = ParagraphStyle(
            "NasaqBody",
            parent=styles["BodyText"],
            fontName=font_name,
            fontSize=10,
            leading=15,
            alignment=TA_RIGHT if written["language"] == "ar" else 0,
        )
        heading = ParagraphStyle(
            "NasaqHeading",
            parent=styles["Heading1"],
            fontName=font_name,
            alignment=TA_RIGHT if written["language"] == "ar" else 0,
        )
        document = SimpleDocTemplate(
            str(path),
            pagesize=A4,
            rightMargin=20 * mm,
            leftMargin=20 * mm,
            topMargin=20 * mm,
            bottomMargin=20 * mm,
            title=written["title"],
            author="Nasaq AI Demo Mode",
        )
        story = [
            Paragraph(escape(written["title"]), heading),
            Spacer(1, 6 * mm),
            Paragraph(
                "DEMO MODE — generated from simulated research; not for factual reliance.", body
            ),
            Spacer(1, 4 * mm),
        ]
        for line in written["content"].splitlines():
            text = line.strip()
            if not text:
                story.append(Spacer(1, 3 * mm))
                continue
            selected_style = heading if text.startswith("#") else body
            cleaned = text.lstrip("#> ")
            if text.startswith("- "):
                cleaned = f"• {text[2:]}"
            story.append(Paragraph(escape(cleaned), selected_style))
        document.build(story)

    @staticmethod
    def _unicode_font() -> str:
        font_name = "NasaqUnicode"
        if font_name in pdfmetrics.getRegisteredFontNames():
            return font_name
        candidates = (
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("C:/Windows/Fonts/arial.ttf"),
            Path("C:/Windows/Fonts/calibri.ttf"),
        )
        for candidate in candidates:
            if candidate.is_file():
                pdfmetrics.registerFont(TTFont(font_name, str(candidate)))
                return font_name
        return "Helvetica"

    @staticmethod
    def _write_docx(path: Path, written: dict[str, Any]) -> None:
        document = Document()
        document.core_properties.title = written["title"]
        document.core_properties.author = "Nasaq AI Demo Mode"
        title = document.add_heading(written["title"], level=0)
        notice = document.add_paragraph(
            "DEMO MODE — generated from simulated research; not for factual reliance."
        )
        if written["language"] == "ar":
            title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            notice.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for line in written["content"].splitlines():
            text = line.strip()
            if not text:
                continue
            if text.startswith("## "):
                paragraph = document.add_heading(text[3:], level=1)
            elif text.startswith("- "):
                paragraph = document.add_paragraph(text[2:], style="List Bullet")
            else:
                paragraph = document.add_paragraph(text.lstrip("> "))
            if written["language"] == "ar":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        document.save(path)

    @staticmethod
    def _artifact(
        path: Path,
        output_format: str,
        context: ExecutionContext,
        node_id: str,
    ) -> AgentArtifact:
        checksum = hashlib.sha256(path.read_bytes()).hexdigest()
        artifact_key = hashlib.sha256(
            f"{context.workflow_run_id}:{node_id}:{output_format}".encode()
        ).hexdigest()[:24]
        return AgentArtifact(
            id=artifact_key,
            owner_id=context.user_id,
            workflow_run_id=context.workflow_run_id,
            node_id=node_id,
            path=str(path.resolve()),
            storage_key=str(path.relative_to(path.parents[2])).replace("\\", "/"),
            mime_type=MIME_TYPES[output_format],
            file_name=path.name,
            file_size=path.stat().st_size,
            checksum_sha256=checksum,
        )
